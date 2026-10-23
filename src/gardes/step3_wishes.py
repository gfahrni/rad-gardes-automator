from __future__ import annotations

import datetime as dt
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from .colors import (
    STATUS_CYAN,
    STATUS_GREEN,
    STATUS_RED,
    is_neutral,
    read_fill,
    status_from_rgb,
)
from .errors import GardesError
from .jsonio import generated_at, sha256_file, write_json
from .xlsx_utils import load_workbook_cached

SHEET = "Feuille 1"
HEADER_ROW = 1
FIRST_DATE_ROW = 6
FIRST_PERSON_COL = 4  # D


def _as_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return None
    return None


def _read_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    duplicates: list[str] = []
    for col in range(FIRST_PERSON_COL, ws.max_column + 1):
        value = ws.cell(HEADER_ROW, col).value
        if value is None:
            continue
        if not isinstance(value, str):
            raise GardesError(
                f"En-tête de personne non textuel en {ws.cell(HEADER_ROW, col).coordinate} : {value!r}.",
                code="wishes_header",
            )
        name = value.strip()
        if not name:
            continue
        if name in headers:
            duplicates.append(name)
        headers[name] = col
    if duplicates:
        raise GardesError(f"Noms en double dans les souhaits : {sorted(set(duplicates))}", code="wishes_duplicate_name")
    return headers


def _read_dates(ws) -> dict[dt.date, int]:
    dates: dict[dt.date, int] = {}
    for row in range(FIRST_DATE_ROW, ws.max_row + 1):
        date = _as_date(ws.cell(row, 1).value)
        if date is None:
            continue
        if date in dates:
            raise GardesError(
                f"Date {date.isoformat()} présente deux fois dans les souhaits.",
                code="wishes_duplicate_date",
            )
        dates[date] = row
    if not dates:
        raise GardesError("Aucune date lue dans les souhaits.", code="wishes_empty")
    return dates


def _base_rgb(ws, row: int) -> str:
    """Fond de base de la ligne : colonne A, puis B/C, sinon blanc."""
    for col in (1, 2, 3):
        rgb, _ = read_fill(ws.cell(row, col))
        if rgb is not None:
            return rgb
    return "FFFFFFFF"


def parse_wishes(
    wishes_path: str | Path,
    calcul: dict,
    planning: dict,
) -> dict:
    wishes_path = Path(wishes_path)
    wb = load_workbook_cached(wishes_path)
    try:
        if SHEET not in wb.sheetnames:
            raise GardesError(f"Feuille '{SHEET}' absente du fichier souhaits.", code="sheet_missing")
        ws = wb[SHEET]
        headers = _read_headers(ws)
        dates = _read_dates(ws)

        known = {p["name"]: p for p in calcul["people"]}
        candidates = set(calcul["candidates"])
        unknown = sorted(set(headers) - set(known))
        if unknown:
            raise GardesError(
                f"Noms inconnus dans les souhaits : {unknown}. La comparaison est exacte "
                "(casse, accents, espaces).",
                code="wishes_unknown_name",
            )
        missing = sorted(candidates - set(headers))
        if missing:
            raise GardesError(
                f"Candidats assignables absents des souhaits : {missing}. "
                "Complète le fichier ou corrige la liste validée.",
                code="wishes_missing_candidate",
            )

        planning_by_date = {_as_date(d["date"]): d for d in planning["days"]}
        wish_dates = set(dates)
        expected_dates = set(planning_by_date)
        if wish_dates != expected_dates:
            missing_dates = sorted(expected_dates - wish_dates)
            extra_dates = sorted(wish_dates - expected_dates)
            raise GardesError(
                f"Dates des souhaits différentes du planning. "
                f"Manquantes : {len(missing_dates)} {missing_dates[:5]}; "
                f"en trop : {len(extra_dates)} {extra_dates[:5]}.",
                code="wishes_dates",
            )

        warnings: list[str] = []
        unexpected: list[dict] = []
        people_out: list[dict] = []
        global_counts: Counter[str] = Counter()

        for name, col in headers.items():
            person = known[name]
            is_candidate = name in candidates
            statuses: list[dict] = []
            vacation_dates: list[str] = []
            for date, row in sorted(dates.items()):
                day = planning_by_date[date]
                category = day["category"]
                cell = ws.cell(row, col)
                rgb, raw = read_fill(cell)
                base = _base_rgb(ws, row)
                status = "neutral"
                if is_neutral(rgb, base):
                    status = "neutral"
                else:
                    classified = status_from_rgb(rgb)
                    if classified is None:
                        unexpected.append(
                            {
                                "cell": cell.coordinate,
                                "person": name,
                                "date": date.isoformat(),
                                "category": category,
                                "rgb": rgb,
                                "raw": raw,
                            }
                        )
                        status = "neutral"
                    elif classified == "vacation":
                        status = "vacation"
                    else:  # volunteer / non_dispo
                        if category in ("weekend", "ferie"):
                            status = classified
                        else:
                            warnings.append(
                                f"Couleur {classified} sur un jour de semaine en {cell.coordinate} "
                                f"({name}, {date.isoformat()}) : ignorée, seul le cyan compte en semaine."
                            )
                            status = "neutral"

                note = None
                if isinstance(cell.value, str) and cell.value.strip():
                    note = cell.value.strip()
                entry = {
                    "date": date.isoformat(),
                    "category": category,
                    "status": status,
                    "raw_color": rgb,
                    "raw_color_desc": raw,
                    "note": note,
                    "cell": cell.coordinate,
                }
                statuses.append(entry)
                global_counts[status] += 1
                if status == "vacation":
                    vacation_dates.append(date.isoformat())

            people_out.append(
                {
                    "name": name,
                    "candidate": is_candidate,
                    "source_column": get_column_letter(col),
                    "statuses": statuses,
                    "vacation_dates": vacation_dates,
                }
            )

        if unexpected:
            lines = [
                f"{u['cell']} : {u['person']}, {u['date']} ({u['category']}), "
                f"couleur {u['raw']}"
                for u in unexpected
            ]
            allowed = (
                "Couleurs autorisées : vert FF00FF00, cyan FF00FFFF, "
                "rouge FFFF0000, ou neutre (fond de ligne, blanc, gris, "
                "sans remplissage, theme0)."
            )
            raise GardesError(
                f"Couleurs non reconnues dans les souhaits "
                f"({len(unexpected)} cellule(s)).\n"
                f"Fichier : {wishes_path}\n"
                + "\n".join(lines)
                + "\n"
                + allowed
                + "\nCorrige ces couleurs dans le fichier des souhaits, "
                "puis relance le programme.",
                code="wishes_color",
                details={"errors": lines + [allowed]},
            )

        return {
            "schema_version": 1,
            "period": dict(calcul["period"]),
            "source": {
                "path": str(wishes_path.resolve()),
                "sha256": sha256_file(wishes_path),
                "sheet": SHEET,
                "generated_at": generated_at(),
            },
            "people": people_out,
            "counts": {
                "people": len(people_out),
                "candidates": len(candidates),
                "dates": len(dates),
                "statuses": dict(global_counts),
            },
            "warnings": warnings,
        }
    finally:
        wb.close()


def write_wishes(data: dict, out: str | Path) -> None:
    write_json(out, data)
