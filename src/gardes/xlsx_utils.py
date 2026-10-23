from __future__ import annotations

import posixpath
import re
import warnings
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .errors import GardesError


def load_workbook_cached(path: str | Path, *, data_only: bool = True):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*Conditional Formatting extension.*")
        return openpyxl.load_workbook(path, data_only=data_only, read_only=False)


def _attr(tag: str, name: str) -> str | None:
    m = re.search(rf'\b{re.escape(name)}="([^"]*)"', tag)
    return m.group(1) if m else None


def sheet_xml_path_from_zip(data: dict[str, bytes], sheet_name: str) -> str:
    workbook = data.get("xl/workbook.xml")
    rels = data.get("xl/_rels/workbook.xml.rels")
    if workbook is None or rels is None:
        raise GardesError("Classeur invalide : workbook.xml ou rels absent.", code="xlsx_invalid")
    wb_text = workbook.decode("utf-8")
    rel_id = None
    for m in re.finditer(r"<sheet\b[^>]*>", wb_text):
        tag = m.group(0)
        if _attr(tag, "name") == sheet_name:
            rel_id = _attr(tag, "r:id")
            break
    if rel_id is None:
        raise GardesError(f"Feuille '{sheet_name}' absente du classeur.", code="sheet_missing")
    rel_text = rels.decode("utf-8")
    target = None
    for m in re.finditer(r"<Relationship\b[^>]*>", rel_text):
        tag = m.group(0)
        if _attr(tag, "Id") == rel_id:
            target = _attr(tag, "Target")
            break
    if target is None:
        raise GardesError(f"Relation '{rel_id}' introuvable dans le classeur.", code="xlsx_invalid")
    if target.startswith("/"):
        path = target.lstrip("/")
    else:
        path = posixpath.normpath(posixpath.join("xl", target))
    if path not in data:
        raise GardesError(f"Feuille XML introuvable : {path}", code="xlsx_invalid")
    return path


def _style_for_cell(text: str, cell: str) -> str | None:
    """Style d'une cellule absente : même colonne avant/après, puis même ligne."""
    m = re.fullmatch(r"([A-Z]+)(\d+)", cell)
    if m is None:
        return None
    col, row = m.group(1), int(m.group(2))
    for r in (row - 1, row + 1):
        probe = re.search(rf'<c r="{col}{r}"([^>]*)>', text)
        if probe is not None:
            style = _attr(probe.group(1), "s")
            if style:
                return style
    col_idx = column_index_from_string(col)
    for delta in (-1, 1):
        neighbour = get_column_letter(col_idx + delta)
        if neighbour == col:
            continue
        probe = re.search(rf'<c r="{neighbour}{row}"([^>]*)>', text)
        if probe is not None:
            style = _attr(probe.group(1), "s")
            if style:
                return style
    return None


def _insert_cell(text: str, cell: str, value: str) -> str:
    m = re.fullmatch(r"([A-Z]+)(\d+)", cell)
    if m is None:
        raise GardesError(f"Référence de cellule invalide : {cell}.", code="cell_invalid")
    col, row = m.group(1), int(m.group(2))
    row_pattern = re.compile(rf'<row r="{row}"[^>]*>.*?</row>', re.S)
    rm = row_pattern.search(text)
    if rm is None:
        raise GardesError(
            f"Ligne {row} absente du XML : écriture de {cell} impossible.", code="cell_missing"
        )
    row_text = rm.group(0)
    style = _style_for_cell(text, cell)
    attrs = f' s="{style}"' if style else ""
    new_cell = f'<c r="{cell}"{attrs} t="inlineStr"><is><t>{escape(value)}</t></is></c>'
    target_idx = column_index_from_string(col)
    insert_pos = None
    for cm in re.finditer(r'<c r="([A-Z]+)\d+"', row_text):
        if column_index_from_string(cm.group(1)) > target_idx:
            insert_pos = cm.start()
            break
    if insert_pos is None:
        insert_pos = row_text.rfind("</row>")
    new_row = row_text[:insert_pos] + new_cell + row_text[insert_pos:]
    return text[: rm.start()] + new_row + text[rm.end() :]


def _patch_cell(text: str, cell: str, value: str) -> str:
    pattern = re.compile(
        r'<c r="' + re.escape(cell) + r'"(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</c>)',
        re.S,
    )
    m = pattern.search(text)
    if m is None:
        return _insert_cell(text, cell, value)
    body = m.group("body")
    if body is not None:
        stripped = body.strip()
        # Une cellule vide peut contenir un cache <v></v> ; tout le reste est un contenu.
        if stripped and not re.fullmatch(r"<v>\s*</v>", stripped):
            raise GardesError(
                f"Cellule {cell} non vide dans le fichier source : elle ne devrait pas être un trou.",
                code="cell_not_empty",
            )
    attrs = m.group("attrs") or ""
    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
    replacement = f'<c r="{cell}"{attrs} t="inlineStr"><is><t>{escape(value)}</t></is></c>'
    return text[: m.start()] + replacement + text[m.end() :]


def _set_full_calc_on_load(workbook_xml: str) -> str:
    m = re.search(r"<calcPr\b[^>]*>", workbook_xml)
    if m is not None:
        tag = m.group(0)
        if "fullCalcOnLoad=" in tag:
            new_tag = re.sub(r'fullCalcOnLoad="[^"]*"', 'fullCalcOnLoad="1"', tag)
        elif tag.endswith("/>"):
            new_tag = tag[:-2] + ' fullCalcOnLoad="1"/>'
        else:
            new_tag = tag[:-1] + ' fullCalcOnLoad="1">'
        return workbook_xml[: m.start()] + new_tag + workbook_xml[m.end() :]
    if "</workbook>" in workbook_xml:
        return workbook_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>', 1)
    raise GardesError("Impossible d'ajouter calcPr : workbook.xml invalide.", code="xlsx_invalid")


def patch_workbook_cells(
    src: str | Path,
    dst: str | Path,
    values: dict[str, str],
    *,
    sheet_name: str = "Feuille1",
) -> None:
    src, dst = Path(src), Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(src) as zin:
        infos = zin.infolist()
        data = {info.filename: zin.read(info.filename) for info in infos}

    sheet_xml = sheet_xml_path_from_zip(data, sheet_name)
    text = data[sheet_xml].decode("utf-8")
    for cell, value in values.items():
        text = _patch_cell(text, cell, value)
    data[sheet_xml] = text.encode("utf-8")

    wb_xml = data.get("xl/workbook.xml")
    if wb_xml is None:
        raise GardesError("Classeur invalide : workbook.xml absent.", code="xlsx_invalid")
    data["xl/workbook.xml"] = _set_full_calc_on_load(wb_xml.decode("utf-8")).encode("utf-8")

    with zipfile.ZipFile(dst, "w") as zout:
        for info in infos:
            zout.writestr(info, data[info.filename])


def verify_patch(
    src: str | Path,
    dst: str | Path,
    *,
    sheet_name: str = "Feuille1",
    expected_values: dict[str, str] | None = None,
    fixed_cells: dict[str, str] | None = None,
) -> None:
    expected_values = expected_values or {}
    fixed_cells = fixed_cells or {}

    with zipfile.ZipFile(dst) as z:
        bad = z.testzip()
        if bad is not None:
            raise GardesError(f"Archive XLSX corrompue : {bad}", code="xlsx_corrupt")
        dst_data = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(src) as z:
        src_data = {n: z.read(n) for n in z.namelist()}

    src_xml = sheet_xml_path_from_zip(src_data, sheet_name)
    dst_xml = sheet_xml_path_from_zip(dst_data, sheet_name)
    src_text = src_data[src_xml].decode("utf-8")
    dst_text = dst_data[dst_xml].decode("utf-8")
    if "<conditionalFormatting" in src_text and "<conditionalFormatting" not in dst_text:
        raise GardesError(
            "La mise en forme conditionnelle a été perdue pendant l'écriture.",
            code="xlsx_cf_lost",
        )

    wb_src = load_workbook_cached(src)
    wb_dst = load_workbook_cached(dst)
    try:
        ws_src = wb_src[sheet_name]
        ws_dst = wb_dst[sheet_name]
        for cell, value in expected_values.items():
            got = ws_dst[cell].value
            if got != value:
                raise GardesError(
                    f"Écriture incohérente en {cell} : attendu {value!r}, obtenu {got!r}.",
                    code="xlsx_verify",
                )
        for cell, value in fixed_cells.items():
            got = ws_dst[cell].value
            if got != value:
                raise GardesError(
                    f"Cellule fixe {cell} modifiée : attendu {value!r}, obtenu {got!r}.",
                    code="xlsx_fixed_changed",
                )
        # Vérifie que tous les caches de formules et les littéraux non patchés sont préservés.
        for row in ws_src.iter_rows():
            for cell in row:
                if cell.coordinate in expected_values:
                    continue
                old = cell.value
                if old is None:
                    continue
                new = ws_dst[cell.coordinate].value
                if new != old:
                    raise GardesError(
                        f"Valeur perdue ou modifiée en {cell.coordinate} : {old!r} -> {new!r}.",
                        code="xlsx_cache_lost",
                    )
    finally:
        wb_src.close()
        wb_dst.close()
