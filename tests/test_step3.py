from __future__ import annotations

import copy
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from gardes.errors import GardesError
from gardes.step1_pdf import parse_pdf, validate_calcul
from gardes.step2_planning import parse_planning
from gardes.step3_wishes import parse_wishes, write_wishes


@pytest.fixture()
def base_data(pdf: Path, planning_xlsx: Path):
    calcul = validate_calcul(parse_pdf(pdf, planning_path=planning_xlsx))
    planning = parse_planning(planning_xlsx)
    return calcul, planning


def test_counts_and_statuses(wishes_xlsx: Path, base_data) -> None:
    calcul, planning = base_data
    data = parse_wishes(wishes_xlsx, calcul, planning)
    assert data["counts"]["people"] == 14
    assert data["counts"]["candidates"] == 12
    assert data["counts"]["dates"] == 181
    assert data["counts"]["statuses"]["volunteer"] == 227
    assert data["counts"]["statuses"]["vacation"] == 384
    assert data["counts"]["statuses"]["non_dispo"] == 106
    assert len(data["warnings"]) == 3


def test_specific_cells(wishes_xlsx: Path, base_data) -> None:
    calcul, planning = base_data
    data = parse_wishes(wishes_xlsx, calcul, planning)
    people = {p["name"]: p for p in data["people"]}
    lefebvre = {s["date"]: s for s in people["Lefebvre"]["statuses"]}
    assert lefebvre["2026-11-07"]["status"] == "volunteer"
    assert lefebvre["2026-11-07"]["raw_color"] == "FF00FF00"
    assert lefebvre["2026-12-24"]["status"] == "vacation"
    assert lefebvre["2026-11-27"]["status"] == "neutral"  # rouge en semaine ignoré
    traore = {s["date"]: s for s in people["Traoré"]["statuses"]}
    assert traore["2026-11-14"]["status"] == "neutral"
    assert traore["2026-11-14"]["raw_color_desc"] == "theme0"
    assert traore["2026-11-07"]["status"] == "neutral"
    assert traore["2026-11-07"]["raw_color"] is None
    hoffmann = {s["date"]: s for s in people["Hoffmann"]["statuses"]}
    assert hoffmann["2026-12-24"]["status"] == "neutral"
    assert hoffmann["2026-12-24"]["raw_color"] == "FFFCE5CD"
    assert hoffmann["2027-01-01"]["status"] == "neutral"
    assert hoffmann["2027-01-01"]["raw_color"] == "FFFCE5CD"


def test_warnings_include_only_weekday_reds(wishes_xlsx: Path, base_data) -> None:
    calcul, planning = base_data
    data = parse_wishes(wishes_xlsx, calcul, planning)
    joined = "\n".join(data["warnings"])
    assert "D32" in joined
    assert "D102" in joined
    assert "D137" in joined
    assert "Q67" not in joined
    assert "FFE6B8AF" not in joined


def test_missing_candidate_blocks(wishes_xlsx: Path, base_data) -> None:
    calcul, planning = base_data
    broken = copy.deepcopy(calcul)
    broken["people"].append({"name": "Absent"})
    broken["candidates"].append("Absent")
    with pytest.raises(GardesError, match="absents des souhaits"):
        parse_wishes(wishes_xlsx, broken, planning)


def test_unknown_header_blocks(wishes_xlsx: Path, base_data, tmp_path: Path) -> None:
    calcul, planning = base_data
    wb = openpyxl.load_workbook(wishes_xlsx)
    ws = wb["Feuille 1"]
    ws["D1"] = "Inconnu"
    broken = tmp_path / "broken-wishes.xlsx"
    wb.save(broken)
    wb.close()
    with pytest.raises(GardesError, match="Noms inconnus"):
        parse_wishes(broken, calcul, planning)


def test_unexpected_color_on_candidate_blocks(
    wishes_xlsx: Path, base_data, tmp_path: Path
) -> None:
    calcul, planning = base_data
    wb = openpyxl.load_workbook(wishes_xlsx)
    ws = wb["Feuille 1"]
    ws["D12"].fill = PatternFill(fill_type="solid", start_color="FF123456", end_color="FF123456")
    broken = tmp_path / "broken-color.xlsx"
    wb.save(broken)
    wb.close()
    with pytest.raises(GardesError, match="Couleurs non reconnues") as exc_info:
        parse_wishes(broken, calcul, planning)
    message = str(exc_info.value)
    assert "D12" in message
    assert "FF123456" in message
    assert "Corrige ces couleurs" in message


def test_unexpected_color_on_non_candidate_blocks(
    wishes_xlsx: Path, base_data, tmp_path: Path
) -> None:
    calcul, planning = base_data
    wb = openpyxl.load_workbook(wishes_xlsx)
    ws = wb["Feuille 1"]
    for coord in ("Q67", "Q68"):
        ws[coord].fill = PatternFill(
            fill_type="solid", start_color="FFE6B8AF", end_color="FFE6B8AF"
        )
    broken = tmp_path / "broken-noncandidate.xlsx"
    wb.save(broken)
    wb.close()
    with pytest.raises(GardesError, match="Couleurs non reconnues") as exc_info:
        parse_wishes(broken, calcul, planning)
    message = str(exc_info.value)
    assert "Q67" in message
    assert "Q68" in message
    assert "Hoffmann" in message
    assert "FFE6B8AF" in message
    assert "2 cellule(s)" in message


def test_roundtrip_json(wishes_xlsx: Path, base_data, tmp_path: Path) -> None:
    calcul, planning = base_data
    data = parse_wishes(wishes_xlsx, calcul, planning)
    out = tmp_path / "souhaits-gardes.json"
    write_wishes(data, out)
    loaded = json.loads(out.read_text(encoding="utf-8"))
    assert loaded["counts"]["people"] == 14
    assert loaded["counts"]["dates"] == 181
