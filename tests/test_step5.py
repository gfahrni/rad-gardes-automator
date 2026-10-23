from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from gardes.errors import GardesError
from gardes.step1_pdf import parse_pdf, validate_calcul
from gardes.step2_planning import parse_planning
from gardes.step5_excel import write_excel
from gardes.step5_report import render_report


@pytest.fixture()
def data(pdf: Path, planning_xlsx: Path):
    calcul = validate_calcul(parse_pdf(pdf, planning_path=planning_xlsx))
    planning = parse_planning(planning_xlsx)
    return calcul, planning


def _attribution(person: str = "Lefebvre") -> dict:
    return {
        "period": {"start": "2026-11-01", "end": "2027-04-30"},
        "inputs": {"planning": {"path": "exemple-planning-vide.xlsx"}},
        "solver": {"name": "test", "version": "0", "status": "OPTIMAL"},
        "objective": {"coverage": 1, "red_forced": 0, "green": 1},
        "assignments": [
            {
                "hole_id": "2026-11-07_day",
                "date": "2026-11-07",
                "slot": "day",
                "category": "weekend",
                "person": person,
                "forced_red": False,
                "cell": "F11",
            }
        ],
        "empty_holes": [],
        "quotas": [],
        "red_forced": [],
        "counts": {
            "holes": 88,
            "filled": 1,
            "empty": 87,
            "fixed": 24,
            "weekend_holes": 80,
            "ferie_holes": 8,
        },
        "warnings": [],
    }


def test_write_one_assignment(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    out = tmp_path / "planning-rempli.xlsx"
    written = write_excel(
        planning_xlsx=planning_xlsx,
        planning=planning,
        attribution=_attribution(),
        out_xlsx=out,
        calcul=calcul,
    )
    assert written == {"F11": "Lefebvre"}
    import openpyxl

    wb = openpyxl.load_workbook(out, data_only=True)
    try:
        ws = wb["Feuille1"]
        assert ws["F11"].value == "Lefebvre"
        assert ws["G11"].value == "Renard"  # fixe préservé
        assert ws["I12"].value == "dim."  # cache de formule préservé
        assert ws["AA6"].value == "fer."
    finally:
        wb.close()


def test_insert_missing_cell(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    attribution = _attribution("Nassar")
    attribution["assignments"][0].update(
        {
            "hole_id": "2027-01-09_evening",
            "date": "2027-01-09",
            "slot": "evening",
            "category": "weekend",
            "cell": "Y13",
        }
    )
    out = tmp_path / "planning-rempli.xlsx"
    written = write_excel(
        planning_xlsx=planning_xlsx,
        planning=planning,
        attribution=attribution,
        out_xlsx=out,
        calcul=calcul,
    )
    assert written == {"Y13": "Nassar"}
    import openpyxl

    wb = openpyxl.load_workbook(out, data_only=True)
    try:
        assert wb["Feuille1"]["Y13"].value == "Nassar"
    finally:
        wb.close()


def test_conditional_formatting_and_calcpr_preserved(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    out = tmp_path / "planning-rempli.xlsx"
    write_excel(
        planning_xlsx=planning_xlsx,
        planning=planning,
        attribution=_attribution(),
        out_xlsx=out,
        calcul=calcul,
    )
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        workbook = z.read("xl/workbook.xml").decode("utf-8")
    assert "<conditionalFormatting" in sheet
    assert 'fullCalcOnLoad="1"' in workbook


def test_fixed_cell_is_refused(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    attribution = _attribution()
    attribution["assignments"][0]["cell"] = "G11"  # fixe
    with pytest.raises(GardesError, match="cellule fixe"):
        write_excel(
            planning_xlsx=planning_xlsx,
            planning=planning,
            attribution=attribution,
            out_xlsx=tmp_path / "out.xlsx",
            calcul=calcul,
        )


def test_unknown_person_is_refused(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    with pytest.raises(GardesError, match="absent de la liste"):
        write_excel(
            planning_xlsx=planning_xlsx,
            planning=planning,
            attribution=_attribution("Inconnu"),
            out_xlsx=tmp_path / "out.xlsx",
            calcul=calcul,
        )


def test_in_place_is_refused(data, planning_xlsx: Path, tmp_path: Path) -> None:
    calcul, planning = data
    with pytest.raises(GardesError, match="différent"):
        write_excel(
            planning_xlsx=planning_xlsx,
            planning=planning,
            attribution=_attribution(),
            out_xlsx=planning_xlsx,
            calcul=calcul,
        )


def test_report_contains_key_sections(data) -> None:
    calcul, planning = data
    text = render_report(planning=planning, attribution=_attribution(), calcul=calcul)
    assert "RAPPORT DE GARDES" in text
    assert "Lefebvre" in text
    assert "F11" in text
    assert "TROUS NON POURVUS" in text
    assert "PRÉ-REMPLIS NON TOUCHÉS" in text
    assert "Moreau" in text
