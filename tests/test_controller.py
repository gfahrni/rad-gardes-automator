from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from gardes_gui.controller import PipelineController
from gardes_gui.state import PipelineState


def test_full_pipeline(pdf: Path, planning_xlsx: Path, wishes_xlsx: Path, tmp_path: Path) -> None:
    controller = PipelineController(PipelineState())
    controller.set_inputs(pdf, planning_xlsx, wishes_xlsx, tmp_path / "run")
    attribution = controller.run_auto()
    assert attribution["objective"] == {"coverage": 87, "red_forced": 3, "green": 74}
    assert attribution["counts"]["holes"] == 88
    assert attribution["counts"]["empty"] == 1
    for name in (
        "calcul-gardes.json",
        "gardes-a-remplir.json",
        "souhaits-gardes.json",
        "attribution-gardes.json",
    ):
        assert (tmp_path / "run" / name).exists()
    out_xlsx, report = controller.export()
    assert out_xlsx.exists()
    assert report.exists()
    assert "RAPPORT DE GARDES" in report.read_text(encoding="utf-8")

    wb = openpyxl.load_workbook(out_xlsx, data_only=True)
    try:
        ws = wb["Feuille1"]
        assert ws["G11"].value == "Renard"
        assert ws["I12"].value == "dim."
        for assignment in attribution["assignments"]:
            assert ws[assignment["cell"]].value == assignment["person"]
    finally:
        wb.close()


def test_validate_people_updates_candidates(
    pdf: Path, planning_xlsx: Path, wishes_xlsx: Path, tmp_path: Path
) -> None:
    controller = PipelineController(PipelineState())
    controller.set_inputs(pdf, planning_xlsx, wishes_xlsx, tmp_path / "run")
    controller.scan_planning()
    calcul = controller.build_calcul()
    people = [dict(p) for p in calcul["people"]]
    jost = next(p for p in people if p["name"] == "Jost")
    jost["quotas"]["samedi"] = 1
    calcul = controller.validate_people(people)
    assert "Jost" in calcul["candidates"]
    assert calcul["validated"] is True


def test_inputs_are_required(tmp_path: Path) -> None:
    controller = PipelineController()
    try:
        controller.scan_planning()
    except Exception as exc:  # noqa: BLE001
        assert "Choisis" in str(exc) or "fichiers" in str(exc).lower()
    else:
        raise AssertionError("Une erreur était attendue")
