from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl

from gardes.cli import main


def test_run_auto_validate(pdf: Path, planning_xlsx: Path, wishes_xlsx: Path, tmp_path: Path) -> None:
    workdir = tmp_path / "out"
    code = main(
        [
            "run",
            "--pdf",
            str(pdf),
            "--planning",
            str(planning_xlsx),
            "--wishes",
            str(wishes_xlsx),
            "--workdir",
            str(workdir),
            "--auto-validate",
        ]
    )
    assert code == 0
    for name in (
        "calcul-gardes.json",
        "gardes-a-remplir.json",
        "souhaits-gardes.json",
        "attribution-gardes.json",
        "planning-rempli.xlsx",
        "rapport.txt",
    ):
        assert (workdir / name).exists(), name
    attribution = json.loads((workdir / "attribution-gardes.json").read_text(encoding="utf-8"))
    assert attribution["objective"] == {"coverage": 87, "red_forced": 3, "green": 74}
    assert attribution["counts"]["empty"] == 1
    wb = openpyxl.load_workbook(workdir / "planning-rempli.xlsx", data_only=True)
    try:
        ws = wb["Feuille1"]
        assert ws["G11"].value == "Renard"
        assert ws["I12"].value == "dim."
        assignment_cells = {a["cell"]: a["person"] for a in attribution["assignments"]}
        assert all(ws[cell].value == person for cell, person in assignment_cells.items())
        for slot in json.loads(
            (workdir / "gardes-a-remplir.json").read_text(encoding="utf-8")
        )["slots"]:
            if slot["status"] == "fixed":
                assert ws[slot["cell"]].value == slot["fixed_name"]
    finally:
        wb.close()
    with zipfile.ZipFile(workdir / "planning-rempli.xlsx") as z:
        assert z.testzip() is None
        assert "<conditionalFormatting" in z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode("utf-8")
    report = (workdir / "rapport.txt").read_text(encoding="utf-8")
    assert "RAPPORT DE GARDES" in report
    assert "Objectif : couverture 87/88" in report
    assert (
        main(
            [
                "validate-attribution",
                "--calcul",
                str(workdir / "calcul-gardes.json"),
                "--planning",
                str(workdir / "gardes-a-remplir.json"),
                "--souhaits",
                str(workdir / "souhaits-gardes.json"),
                "--attribution",
                str(workdir / "attribution-gardes.json"),
            ]
        )
        == 0
    )


def test_run_requires_human_validation(
    pdf: Path, planning_xlsx: Path, wishes_xlsx: Path, tmp_path: Path
) -> None:
    workdir = tmp_path / "out"
    args = [
        "run",
        "--pdf",
        str(pdf),
        "--planning",
        str(planning_xlsx),
        "--wishes",
        str(wishes_xlsx),
        "--workdir",
        str(workdir),
    ]
    assert main(args) == 0
    calcul_path = workdir / "calcul-gardes.json"
    calcul = json.loads(calcul_path.read_text(encoding="utf-8"))
    assert calcul["validated"] is False
    assert not (workdir / "attribution-gardes.json").exists()

    assert main(["step1-validate", "--json", str(calcul_path)]) == 0
    calcul = json.loads(calcul_path.read_text(encoding="utf-8"))
    assert calcul["validated"] is True

    assert main(args + ["--calcul", str(calcul_path)]) == 0
    assert (workdir / "attribution-gardes.json").exists()


def test_cli_error_is_clean(tmp_path: Path, capsys) -> None:
    code = main(
        ["step2", "--planning", str(tmp_path / "absent.xlsx"), "--out", str(tmp_path / "out.json")]
    )
    assert code == 1
    captured = capsys.readouterr()
    assert "ERREUR" in captured.err
    assert "introuvable" in captured.err
